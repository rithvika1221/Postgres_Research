#!/usr/bin/env python3
"""
Subscriber Replication Monitor — 4 Users, Variable Throughput
Experiment: 4 users (-c 4 -j 4) with TPS stepping every 2 minutes

Monitors PostgreSQL logical replication metrics on the SUBSCRIBER VM
and writes all data to a single .xlsx file with 4 sheets.

Each row includes "TPS Phase" and "Target TPS" columns.

TPS Schedule (must match load_generator.py):
  Phase 1 (0:00 - 2:00)  →  100 TPS
  Phase 2 (2:00 - 4:00)  →  500 TPS
  Phase 3 (4:00 - 6:00)  → 1000 TPS
  Phase 4 (6:00 - 8:00)  → 2500 TPS
  Phase 5 (8:00 - 10:00) → 5000 TPS

Sheets:
  1. Subscription_Lag   — Subscription status, LSN lag, worker info
  2. Apply_Stats        — Apply worker activity and lag
  3. Worker_Activity    — Active worker processes, wait events
  4. Table_Status       — Per-table subscription sync state

Run on: SUBSCRIBER VM
Requires: psycopg2, openpyxl
Usage:   python3 monitor_subscriber_xlsx.py
"""

import os
import sys
import time
import argparse
from datetime import datetime

try:
    import psycopg2
    import psycopg2.extras
except ImportError:
    print("ERROR: psycopg2 not installed. Run: pip install psycopg2-binary")
    sys.exit(1)

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("ERROR: openpyxl not installed. Run: pip install openpyxl")
    sys.exit(1)


# ─── Configuration ───────────────────────────────────────────────────────────
DB_CONFIG = {
    "host": "localhost",
    "port": 5432,
    "dbname": "sub",
    "user": "postgres",
    "password": "Aarush@123",
}

PUBLISHER_CONFIG = {
    "host": "172.16.0.4",
    "port": 5432,
    "dbname": "pub",
    "user": "postgres",
    "password": "Aarush@123",
}

SAMPLE_INTERVAL = 5
DURATION_SECONDS = 660
OUTPUT_DIR = "."

TPS_SCHEDULE = [
    ("Phase 1: 100 TPS",   100,  120),
    ("Phase 2: 500 TPS",   500,  120),
    ("Phase 3: 1000 TPS", 1000,  120),
    ("Phase 4: 2500 TPS", 2500,  120),
    ("Phase 5: 5000 TPS", 5000,  120),
]


# ─── TPS Phase Detection ────────────────────────────────────────────────────

def get_tps_phase_from_time(elapsed_seconds):
    cumulative = 0
    for phase_name, target_tps, duration in TPS_SCHEDULE:
        if elapsed_seconds < cumulative + duration:
            return phase_name, target_tps
        cumulative += duration
    return "Completed", 0


# ─── Publisher LSN Fetch ─────────────────────────────────────────────────────

def get_publisher_lsn():
    try:
        conn = psycopg2.connect(**PUBLISHER_CONFIG)
        cur = conn.cursor()
        cur.execute("SELECT pg_current_wal_lsn()::text")
        lsn = cur.fetchone()[0]
        conn.close()
        return lsn
    except Exception:
        return None


def lsn_to_int(lsn_str):
    if not lsn_str or lsn_str == "":
        return 0
    try:
        hi, lo = lsn_str.split('/')
        return (int(hi, 16) << 32) + int(lo, 16)
    except Exception:
        return 0


# ─── Excel Styling ───────────────────────────────────────────────────────────
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
HEADER_FILL = PatternFill(start_color="548235", end_color="548235", fill_type="solid")
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)


def style_header_row(ws, num_cols):
    for col in range(1, num_cols + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGNMENT
        cell.border = THIN_BORDER
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def auto_fit_columns(ws):
    for col in ws.columns:
        max_length = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value is not None:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_length + 4, 40)


# ─── SQL Queries ─────────────────────────────────────────────────────────────

QUERY_SUBSCRIPTION_LAG = """
SELECT
    now() AS sample_time,
    s.subname, s.pid, s.received_lsn::text, s.latest_end_lsn::text,
    s.last_msg_send_time::text, s.last_msg_receipt_time::text,
    EXTRACT(EPOCH FROM (now() - s.last_msg_receipt_time))::numeric(10,3),
    (SELECT count(*) FROM ingest_data)
FROM pg_stat_subscription s WHERE s.subname = 'mysub';
"""

QUERY_APPLY_STATS = """
SELECT
    now() AS sample_time,
    s.subname, s.pid, s.received_lsn::text, s.latest_end_lsn::text,
    s.last_msg_send_time::text, s.last_msg_receipt_time::text,
    a.state, a.wait_event_type, a.wait_event, a.query,
    EXTRACT(EPOCH FROM (now() - a.query_start))::numeric(10,3)
FROM pg_stat_subscription s
LEFT JOIN pg_stat_activity a ON a.pid = s.pid
WHERE s.subname = 'mysub';
"""

QUERY_WORKER_ACTIVITY = """
SELECT
    now() AS sample_time,
    pid, usename, application_name, state,
    wait_event_type, wait_event, backend_type, query,
    EXTRACT(EPOCH FROM (now() - query_start))::numeric(10,3)
FROM pg_stat_activity
WHERE datname = 'sub'
  AND backend_type IN ('logical replication worker', 'client backend')
ORDER BY backend_type, pid;
"""

QUERY_TABLE_STATUS = """
SELECT
    now() AS sample_time,
    sr.srsubid, s.subname, sr.srrelid::regclass::text,
    sr.srsubstate, sr.srsublsn::text
FROM pg_subscription_rel sr
JOIN pg_subscription s ON s.oid = sr.srsubid;
"""


# ─── Sheet Definitions ───────────────────────────────────────────────────────

SHEETS = [
    {
        "name": "Subscription_Lag",
        "query": QUERY_SUBSCRIPTION_LAG,
        "headers": [
            "Sample Time", "TPS Phase", "Target TPS",
            "Subscription", "Worker PID",
            "Received LSN", "Latest End LSN",
            "Last Msg Send Time", "Last Msg Receipt Time",
            "Seconds Since Last Msg", "Subscriber Row Count"
        ],
    },
    {
        "name": "Apply_Stats",
        "query": QUERY_APPLY_STATS,
        "headers": [
            "Sample Time", "TPS Phase", "Target TPS",
            "Subscription", "PID",
            "Received LSN", "Latest End LSN",
            "Last Msg Send Time", "Last Msg Receipt Time",
            "Worker State", "Wait Event Type", "Wait Event",
            "Current Query", "Query Runtime (sec)"
        ],
    },
    {
        "name": "Worker_Activity",
        "query": QUERY_WORKER_ACTIVITY,
        "headers": [
            "Sample Time", "TPS Phase", "Target TPS",
            "PID", "User", "Application",
            "State", "Wait Event Type", "Wait Event",
            "Backend Type", "Query", "Query Runtime (sec)"
        ],
    },
    {
        "name": "Table_Status",
        "query": QUERY_TABLE_STATUS,
        "headers": [
            "Sample Time", "TPS Phase", "Target TPS",
            "Subscription OID", "Subscription Name",
            "Table Name", "Sync State", "Sync LSN"
        ],
    },
]


# ─── Main Monitor ────────────────────────────────────────────────────────────

def create_workbook():
    wb = Workbook()
    wb.remove(wb.active)
    for sheet_def in SHEETS:
        ws = wb.create_sheet(title=sheet_def["name"])
        ws.append(sheet_def["headers"])
        style_header_row(ws, len(sheet_def["headers"]))
    return wb


def collect_sample(conn, sheet_data, elapsed_seconds):
    cur = conn.cursor()
    phase_name, target_tps = get_tps_phase_from_time(elapsed_seconds)

    for i, sheet_def in enumerate(SHEETS):
        try:
            cur.execute(sheet_def["query"])
            rows = cur.fetchall()
            for row in rows:
                processed = []
                first = True
                for val in row:
                    if first:
                        if isinstance(val, datetime):
                            processed.append(val.strftime("%Y-%m-%d %H:%M:%S"))
                        else:
                            processed.append(str(val))
                        processed.append(phase_name)
                        processed.append(target_tps)
                        first = False
                        continue
                    if val is None:
                        processed.append("")
                    elif isinstance(val, datetime):
                        processed.append(val.strftime("%Y-%m-%d %H:%M:%S"))
                    elif isinstance(val, bool):
                        processed.append("Yes" if val else "No")
                    else:
                        processed.append(val)
                sheet_data[i].append(processed)

            if not rows:
                no_data = [datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                           phase_name, target_tps, "NO DATA"]
                no_data.extend([""] * (len(sheet_def["headers"]) - 4))
                sheet_data[i].append(no_data)

        except Exception as e:
            error_row = [datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                         phase_name, target_tps, f"QUERY ERROR: {e}"]
            error_row.extend([""] * (len(sheet_def["headers"]) - 4))
            sheet_data[i].append(error_row)
            conn.rollback()

    conn.commit()
    return phase_name, target_tps


def save_workbook(wb, sheet_data, output_path):
    for i, sheet_def in enumerate(SHEETS):
        ws = wb[sheet_def["name"]]
        for row in sheet_data[i]:
            ws.append(row)
        auto_fit_columns(ws)
        if ws.max_row > 1:
            ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    wb.save(output_path)


def main():
    parser = argparse.ArgumentParser(
        description="Subscriber monitor for 4-user variable TPS experiment"
    )
    parser.add_argument("--duration", type=int, default=DURATION_SECONDS)
    parser.add_argument("--interval", type=int, default=SAMPLE_INTERVAL)
    parser.add_argument("--output-dir", type=str, default=OUTPUT_DIR)
    parser.add_argument("--host", type=str, default=DB_CONFIG["host"])
    parser.add_argument("--port", type=int, default=DB_CONFIG["port"])
    parser.add_argument("--no-publisher-check", action="store_true")
    args = parser.parse_args()

    DB_CONFIG["host"] = args.host
    DB_CONFIG["port"] = args.port

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_filename = f"subscriber_metrics_4users_varTPS_{timestamp}.xlsx"
    output_path = os.path.join(args.output_dir, output_filename)

    print("=" * 65)
    print("  Subscriber Replication Monitor — 4 Users, Variable TPS")
    print("  Experiment: 4 users (-c 4 -j 4), TPS stepping every 2 min")
    print("=" * 65)
    print(f"  Host:       {DB_CONFIG['host']}:{DB_CONFIG['port']}")
    print(f"  Database:   {DB_CONFIG['dbname']}")
    print(f"  Publisher:  {PUBLISHER_CONFIG['host']}:{PUBLISHER_CONFIG['port']}")
    print(f"  Interval:   {args.interval}s | Duration: {args.duration}s")
    print(f"  Output:     {output_path}")
    print(f"\n  TPS Schedule:")
    cumulative = 0
    for name, tps, dur in TPS_SCHEDULE:
        m1, s1 = divmod(cumulative, 60)
        m2, s2 = divmod(cumulative + dur, 60)
        print(f"    {m1}:{s1:02d} - {m2}:{s2:02d}  →  {tps:>5,} TPS")
        cumulative += dur
    print("=" * 65)

    print("\n[1/3] Connecting to subscriber database...")
    try:
        conn = psycopg2.connect(**DB_CONFIG)
        conn.autocommit = False
        print("  Connected.")
    except Exception as e:
        print(f"  FAILED: {e}")
        sys.exit(1)

    print("[2/3] Verifying subscription 'mysub'...")
    try:
        cur = conn.cursor()
        cur.execute("SELECT subname FROM pg_subscription WHERE subname = 'mysub'")
        sub = cur.fetchone()
        print(f"  Subscription '{sub[0]}' exists." if sub else "  WARNING: not found!")
        conn.commit()
    except Exception as e:
        print(f"  WARNING: {e}")
        conn.rollback()

    if not args.no_publisher_check:
        print("[3/3] Testing publisher connectivity...")
        pub_lsn = get_publisher_lsn()
        if pub_lsn:
            print(f"  Publisher reachable. WAL LSN: {pub_lsn}")
        else:
            print("  WARNING: Cannot reach publisher.")
    else:
        print("[3/3] Skipping publisher check.")

    wb = create_workbook()
    sheet_data = [[] for _ in SHEETS]

    print(f"\n  Monitoring started at {datetime.now().strftime('%H:%M:%S')}...\n")

    start_time = time.time()
    end_time = start_time + args.duration
    sample_count = 0

    try:
        while time.time() < end_time:
            sample_start = time.time()
            elapsed = time.time() - start_time

            phase_name, target_tps = collect_sample(conn, sheet_data, elapsed)
            sample_count += 1

            try:
                cur = conn.cursor()
                cur.execute("SELECT count(*) FROM ingest_data")
                row_count = cur.fetchone()[0]
                conn.commit()
            except Exception:
                row_count = "N/A"
                conn.rollback()

            lag_str = "N/A"
            if not args.no_publisher_check:
                pub_lsn = get_publisher_lsn()
                if pub_lsn:
                    try:
                        cur = conn.cursor()
                        cur.execute("""
                            SELECT latest_end_lsn::text
                            FROM pg_stat_subscription
                            WHERE subname = 'mysub' LIMIT 1
                        """)
                        result = cur.fetchone()
                        conn.commit()
                        if result and result[0]:
                            lag_bytes = lsn_to_int(pub_lsn) - lsn_to_int(result[0])
                            lag_str = f"{lag_bytes:,} bytes"
                    except Exception:
                        conn.rollback()

            elapsed_int = int(elapsed)
            remaining = max(0, int(end_time - time.time()))
            print(
                f"  Sample {sample_count:>4} | "
                f"{elapsed_int:>4}s | "
                f"{phase_name:25s} | "
                f"Rows: {row_count:>10} | "
                f"Lag: {lag_str:>15s} | "
                f"Left: {remaining}s"
            )

            elapsed_sample = time.time() - sample_start
            sleep_time = max(0, args.interval - elapsed_sample)
            if sleep_time > 0:
                time.sleep(sleep_time)

    except KeyboardInterrupt:
        print("\n  Ctrl+C — stopping monitor...")

    finally:
        print(f"\n  Saving {output_filename}...")
        save_workbook(wb, sheet_data, output_path)
        conn.close()

        total_rows = sum(len(sd) for sd in sheet_data)
        print("\n" + "=" * 65)
        print("  MONITORING COMPLETE")
        print("=" * 65)
        print(f"  Samples:    {sample_count}")
        print(f"  Total rows: {total_rows} across 4 sheets")
        print(f"  Output:     {output_path}")
        print(f"  File size:  {os.path.getsize(output_path):,} bytes")
        print("=" * 65)


if __name__ == "__main__":
    main()
