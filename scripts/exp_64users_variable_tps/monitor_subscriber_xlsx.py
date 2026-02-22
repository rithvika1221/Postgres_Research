#!/usr/bin/env python3
"""
Subscriber Replication Monitor — 64 users (-c 64 -j 16), Variable Throughput
Each row includes "TPS Phase" and "Target TPS" columns.

TPS Schedule: 100 -> 500 -> 1000 -> 2500 -> 5000 (2 min each)

Run on: SUBSCRIBER VM
Requires: psycopg2, openpyxl
"""

import os, sys, time, argparse
from datetime import datetime

try:
    import psycopg2, psycopg2.extras
except ImportError:
    print("ERROR: pip install psycopg2-binary"); sys.exit(1)
try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("ERROR: pip install openpyxl"); sys.exit(1)

DB_CONFIG = {"host": "localhost", "port": 5432, "dbname": "sub", "user": "postgres", "password": "Aarush@123"}
PUB_CONFIG = {"host": "172.16.0.4", "port": 5432, "dbname": "pub", "user": "postgres", "password": "Aarush@123"}
SAMPLE_INTERVAL = 5
DURATION_SECONDS = 660
OUTPUT_DIR = "."

TPS_SCHEDULE = [
    ("Phase 1: 100 TPS",   100,  120),
    ("Phase 2: 500 TPS",   500,  120),
    ("Phase 3: 1000 TPS", 1000,  120),
    ("Phase 4: 2500 TPS", 2500,  120),
    ("Phase 5: 5000 TPS", 5000,  120),
]

def get_tps_phase(elapsed):
    cumulative = 0
    for name, tps, dur in TPS_SCHEDULE:
        if elapsed < cumulative + dur: return name, tps
        cumulative += dur
    return "Completed", 0

def get_publisher_lsn():
    try:
        c = psycopg2.connect(**PUB_CONFIG); cur = c.cursor()
        cur.execute("SELECT pg_current_wal_lsn()::text"); lsn = cur.fetchone()[0]; c.close(); return lsn
    except Exception: return None

def lsn_to_int(s):
    if not s: return 0
    try:
        hi, lo = s.split("/"); return (int(hi, 16) << 32) + int(lo, 16)
    except Exception: return 0

HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
HEADER_FILL = PatternFill(start_color="548235", end_color="548235", fill_type="solid")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
THIN_BORDER = Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin"))

def style_header_row(ws, n):
    for col in range(1, n + 1):
        c = ws.cell(row=1, column=col)
        c.font = HEADER_FONT; c.fill = HEADER_FILL; c.alignment = HEADER_ALIGN; c.border = THIN_BORDER
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions

def auto_fit_columns(ws):
    for col in ws.columns:
        mx = max((len(str(c.value)) for c in col if c.value is not None), default=0)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(mx + 4, 40)

QUERIES = [
    ("Subscription_Lag",
     "SELECT now(), s.subname, s.pid, s.received_lsn::text, s.latest_end_lsn::text, s.last_msg_send_time::text, s.last_msg_receipt_time::text, EXTRACT(EPOCH FROM (now() - s.last_msg_receipt_time))::numeric(10,3), (SELECT count(*) FROM ingest_data) FROM pg_stat_subscription s WHERE s.subname = 'mysub'",
     ["Sample Time","TPS Phase","Target TPS","Subscription","Worker PID","Received LSN","Latest End LSN","Last Msg Send Time","Last Msg Receipt Time","Seconds Since Last Msg","Subscriber Row Count"]),
    ("Apply_Stats",
     "SELECT now(), s.subname, s.pid, s.received_lsn::text, s.latest_end_lsn::text, s.last_msg_send_time::text, s.last_msg_receipt_time::text, a.state, a.wait_event_type, a.wait_event, a.query, EXTRACT(EPOCH FROM (now() - a.query_start))::numeric(10,3) FROM pg_stat_subscription s LEFT JOIN pg_stat_activity a ON a.pid = s.pid WHERE s.subname = 'mysub'",
     ["Sample Time","TPS Phase","Target TPS","Subscription","PID","Received LSN","Latest End LSN","Last Msg Send Time","Last Msg Receipt Time","Worker State","Wait Event Type","Wait Event","Current Query","Query Runtime (sec)"]),
    ("Worker_Activity",
     "SELECT now(), pid, usename, application_name, state, wait_event_type, wait_event, backend_type, query, EXTRACT(EPOCH FROM (now() - query_start))::numeric(10,3) FROM pg_stat_activity WHERE datname = 'sub' AND backend_type IN ('logical replication worker', 'client backend') ORDER BY backend_type, pid",
     ["Sample Time","TPS Phase","Target TPS","PID","User","Application","State","Wait Event Type","Wait Event","Backend Type","Query","Query Runtime (sec)"]),
    ("Table_Status",
     "SELECT now(), sr.srsubid, s.subname, sr.srrelid::regclass::text, sr.srsubstate, sr.srsublsn::text FROM pg_subscription_rel sr JOIN pg_subscription s ON s.oid = sr.srsubid",
     ["Sample Time","TPS Phase","Target TPS","Subscription OID","Subscription Name","Table Name","Sync State","Sync LSN"]),
]

def create_workbook():
    wb = Workbook(); wb.remove(wb.active)
    for name, _, headers in QUERIES:
        ws = wb.create_sheet(title=name); ws.append(headers); style_header_row(ws, len(headers))
    return wb

def collect_sample(conn, sheet_data, elapsed):
    cur = conn.cursor()
    pn, tt = get_tps_phase(elapsed)
    for i, (_, query, headers) in enumerate(QUERIES):
        try:
            cur.execute(query); rows = cur.fetchall()
            for row in rows:
                processed = []; first = True
                for val in row:
                    if first:
                        processed.append(val.strftime("%Y-%m-%d %H:%M:%S") if isinstance(val, datetime) else str(val))
                        processed.append(pn); processed.append(tt); first = False; continue
                    if val is None: processed.append("")
                    elif isinstance(val, bool): processed.append("Yes" if val else "No")
                    elif isinstance(val, datetime): processed.append(val.strftime("%Y-%m-%d %H:%M:%S"))
                    else: processed.append(val)
                sheet_data[i].append(processed)
            if not rows:
                r = [datetime.now().strftime("%Y-%m-%d %H:%M:%S"), pn, tt, "NO DATA"]
                r.extend([""] * (len(headers) - 4)); sheet_data[i].append(r)
        except Exception as e:
            r = [datetime.now().strftime("%Y-%m-%d %H:%M:%S"), pn, tt, f"ERROR: {e}"]
            r.extend([""] * (len(headers) - 4)); sheet_data[i].append(r); conn.rollback()
    conn.commit()
    return pn, tt

def save_workbook(wb, sheet_data, path):
    for i, (name, _, _) in enumerate(QUERIES):
        ws = wb[name]
        for row in sheet_data[i]: ws.append(row)
        auto_fit_columns(ws)
        if ws.max_row > 1: ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    wb.save(path)

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--duration", type=int, default=DURATION_SECONDS)
    parser.add_argument("--interval", type=int, default=SAMPLE_INTERVAL)
    parser.add_argument("--output-dir", type=str, default=OUTPUT_DIR)
    parser.add_argument("--host", type=str, default=DB_CONFIG["host"])
    parser.add_argument("--port", type=int, default=DB_CONFIG["port"])
    parser.add_argument("--no-publisher-check", action="store_true")
    args = parser.parse_args()
    DB_CONFIG["host"] = args.host; DB_CONFIG["port"] = args.port

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    outfile = f"subscriber_metrics_64users_varTPS_{ts}.xlsx"
    outpath = os.path.join(args.output_dir, outfile)

    print("=" * 65)
    print("  Subscriber Monitor — 64 users (-c 64 -j 16), Variable TPS")
    print("=" * 65)
    print(f"  Host: {DB_CONFIG['host']}:{DB_CONFIG['port']} | DB: {DB_CONFIG['dbname']}")
    print(f"  Publisher: {PUB_CONFIG['host']}:{PUB_CONFIG['port']}")
    print(f"  Interval: {args.interval}s | Duration: {args.duration}s | Output: {outpath}")
    cumulative = 0
    for name, tps, dur in TPS_SCHEDULE:
        m1, s1 = divmod(cumulative, 60); m2, s2 = divmod(cumulative + dur, 60)
        print(f"    {m1}:{s1:02d}-{m2}:{s2:02d} -> {tps:>5,} TPS")
        cumulative += dur
    print("=" * 65)

    print("\n  Connecting to subscriber..."); conn = psycopg2.connect(**DB_CONFIG); conn.autocommit = False; print("  Connected.")

    print("  Checking subscription...")
    try:
        cur = conn.cursor(); cur.execute("SELECT subname FROM pg_subscription WHERE subname = 'mysub'")
        sub = cur.fetchone(); print(f"  Subscription '{sub[0]}' OK." if sub else "  WARNING: not found!"); conn.commit()
    except Exception as e: print(f"  WARNING: {e}"); conn.rollback()

    if not args.no_publisher_check:
        print("  Testing publisher..."); lsn = get_publisher_lsn()
        print(f"  Publisher reachable. LSN: {lsn}" if lsn else "  WARNING: Cannot reach publisher.")

    wb = create_workbook(); sheet_data = [[] for _ in QUERIES]
    print(f"\n  Monitoring started at {datetime.now().strftime('%H:%M:%S')}...\n")

    start = time.time(); end = start + args.duration; count = 0
    try:
        while time.time() < end:
            ss = time.time(); elapsed = time.time() - start
            pn, tt = collect_sample(conn, sheet_data, elapsed); count += 1

            try:
                cur = conn.cursor(); cur.execute("SELECT count(*) FROM ingest_data")
                rc = cur.fetchone()[0]; conn.commit()
            except Exception: rc = "N/A"; conn.rollback()

            lag = "N/A"
            if not args.no_publisher_check:
                plsn = get_publisher_lsn()
                if plsn:
                    try:
                        cur = conn.cursor()
                        cur.execute("SELECT latest_end_lsn::text FROM pg_stat_subscription WHERE subname = 'mysub' LIMIT 1")
                        r = cur.fetchone(); conn.commit()
                        if r and r[0]: lag = f"{lsn_to_int(plsn) - lsn_to_int(r[0]):,} bytes"
                    except Exception: conn.rollback()

            print(f"  Sample {count:>4} | {int(elapsed):>4}s | {pn:25s} | Rows: {rc:>10} | Lag: {lag:>15s} | Left: {max(0, int(end - time.time()))}s")
            sl = max(0, args.interval - (time.time() - ss))
            if sl > 0: time.sleep(sl)
    except KeyboardInterrupt: print("\n  Stopping...")
    finally:
        print(f"\n  Saving {outfile}..."); save_workbook(wb, sheet_data, outpath); conn.close()
        tr = sum(len(s) for s in sheet_data)
        print(f"  Done: {count} samples, {tr} rows, {os.path.getsize(outpath):,} bytes")

if __name__ == "__main__":
    main()
