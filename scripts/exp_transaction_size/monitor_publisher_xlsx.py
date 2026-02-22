#!/usr/bin/env python3
"""
Publisher Replication Monitor — Transaction Size Experiment
Each row includes "TPS Phase", "Target TPS", "Workers", and "Rows/Commit" columns.

Phase Schedule (2 min each, 12 min total):
  Phase 1: 1 user,  100 TPS, 1 row/commit
  Phase 2: 4 users, 500 TPS, 10 rows/commit
  Phase 3: 8 users, 1000 TPS, 50 rows/commit
  Phase 4: 16 users, 2500 TPS, 100 rows/commit
  Phase 5: 32 users, 5000 TPS, 500 rows/commit
  Phase 6: 64 users, 5000 TPS, 1000 rows/commit

Run on: PUBLISHER VM
Requires: psycopg2, openpyxl
"""

import os, sys, time, argparse
from datetime import datetime

try:
    import psycopg2, psycopg2.extras
except ImportError:
    print("ERROR: pip install psycopg2-binary"); sys.exit(1)
try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("ERROR: pip install openpyxl"); sys.exit(1)

DB_CONFIG = {"host": "localhost", "port": 5432, "dbname": "pub", "user": "postgres", "password": "Aarush@123"}
SAMPLE_INTERVAL = 5
DURATION_SECONDS = 780
OUTPUT_DIR = "."
STATUS_FILE = "tps_phase.txt"

TPS_SCHEDULE = [
    ("Phase 1: 1 user, 100 TPS, 1 row/txn",       100,  1,    1, 120),
    ("Phase 2: 4 users, 500 TPS, 10 rows/txn",    500,  4,   10, 120),
    ("Phase 3: 8 users, 1000 TPS, 50 rows/txn",  1000,  8,   50, 120),
    ("Phase 4: 16 users, 2500 TPS, 100 rows/txn", 2500, 16,  100, 120),
    ("Phase 5: 32 users, 5000 TPS, 500 rows/txn", 5000, 32,  500, 120),
    ("Phase 6: 64 users, 5000 TPS, 1000 rows/txn",5000, 64, 1000, 120),
]

def get_tps_phase_from_file():
    try:
        if os.path.exists(STATUS_FILE):
            with open(STATUS_FILE, "r") as f:
                lines = f.readlines()
            phase_name = lines[0].strip() if lines else "Unknown"
            target_tps = 0; workers = 0; rows_commit = 0
            for line in lines:
                if line.startswith("target_tps="): target_tps = int(line.split("=")[1].strip())
                elif line.startswith("num_workers="): workers = int(line.split("=")[1].strip())
                elif line.startswith("rows_per_commit="): rows_commit = int(line.split("=")[1].strip())
            return phase_name, target_tps, workers, rows_commit
    except Exception:
        pass
    return "Unknown", 0, 0, 0

def get_tps_phase_from_time(elapsed):
    cumulative = 0
    for name, tps, w, rpc, dur in TPS_SCHEDULE:
        if elapsed < cumulative + dur:
            return name, tps, w, rpc
        cumulative += dur
    return "Completed", 0, 0, 0

def get_current_tps_phase(elapsed):
    name, tps, w, rpc = get_tps_phase_from_file()
    if name == "Unknown" or tps == 0:
        name, tps, w, rpc = get_tps_phase_from_time(elapsed)
    return name, tps, w, rpc

HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
HEADER_FILL = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
THIN_BORDER = Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin"))

def style_header_row(ws, n):
    for col in range(1, n + 1):
        c = ws.cell(row=1, column=col)
        c.font = HEADER_FONT; c.fill = HEADER_FILL; c.alignment = HEADER_ALIGN; c.border = THIN_BORDER
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions

def auto_fit_columns(ws):
    for col in ws.columns:
        mx = max((len(str(c.value)) for c in col if c.value is not None), default=0)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(mx + 4, 40)

QUERIES = [
    ("Replication_Lag",
     "SELECT now(), pid, usename, application_name, client_addr::text, state, sent_lsn::text, write_lsn::text, flush_lsn::text, replay_lsn::text, pg_wal_lsn_diff(sent_lsn, write_lsn), pg_wal_lsn_diff(sent_lsn, flush_lsn), pg_wal_lsn_diff(sent_lsn, replay_lsn), write_lag::text, flush_lag::text, replay_lag::text, sync_state FROM pg_stat_replication",
     ["Sample Time","TPS Phase","Target TPS","Workers","Rows/Commit","PID","User","Application","Client Addr","State","Sent LSN","Write LSN","Flush LSN","Replay LSN","Write Lag (bytes)","Flush Lag (bytes)","Replay Lag (bytes)","Write Lag (time)","Flush Lag (time)","Replay Lag (time)","Sync State"]),
    ("WAL_Stats",
     "SELECT now(), pg_current_wal_lsn()::text, pg_current_wal_insert_lsn()::text, pg_wal_lsn_diff(pg_current_wal_insert_lsn(), pg_current_wal_lsn()), (SELECT count(*) FROM pg_ls_waldir()), pg_size_pretty(sum(size)) FROM pg_ls_waldir()",
     ["Sample Time","TPS Phase","Target TPS","Workers","Rows/Commit","Current WAL LSN","Current WAL Insert LSN","Insert Ahead (bytes)","WAL File Count","Total WAL Size"]),
    ("Database_Stats",
     "SELECT now(), numbackends, xact_commit, xact_rollback, blks_read, blks_hit, tup_returned, tup_fetched, tup_inserted, tup_updated, tup_deleted, conflicts, temp_files, temp_bytes, deadlocks, blk_read_time, blk_write_time FROM pg_stat_database WHERE datname = 'pub'",
     ["Sample Time","TPS Phase","Target TPS","Workers","Rows/Commit","Backends","Commits","Rollbacks","Blocks Read","Blocks Hit","Tuples Returned","Tuples Fetched","Tuples Inserted","Tuples Updated","Tuples Deleted","Conflicts","Temp Files","Temp Bytes","Deadlocks","Block Read Time (ms)","Block Write Time (ms)"]),
    ("Connections",
     "SELECT now(), state, wait_event_type, wait_event, count(*), string_agg(DISTINCT application_name, ', ') FROM pg_stat_activity WHERE datname = 'pub' GROUP BY state, wait_event_type, wait_event ORDER BY count(*) DESC",
     ["Sample Time","TPS Phase","Target TPS","Workers","Rows/Commit","State","Wait Event Type","Wait Event","Connection Count","Applications"]),
    ("Replication_Slots",
     "SELECT now(), slot_name, plugin, slot_type, active, active_pid, restart_lsn::text, confirmed_flush_lsn::text, pg_wal_lsn_diff(pg_current_wal_lsn(), restart_lsn), pg_size_pretty(pg_wal_lsn_diff(pg_current_wal_lsn(), restart_lsn)), wal_status FROM pg_replication_slots",
     ["Sample Time","TPS Phase","Target TPS","Workers","Rows/Commit","Slot Name","Plugin","Slot Type","Active","Active PID","Restart LSN","Confirmed Flush LSN","Retained WAL (bytes)","Retained WAL (pretty)","WAL Status"]),
]

def create_workbook():
    wb = Workbook(); wb.remove(wb.active)
    for name, _, headers in QUERIES:
        ws = wb.create_sheet(title=name); ws.append(headers); style_header_row(ws, len(headers))
    return wb

def collect_sample(conn, sheet_data, elapsed):
    cur = conn.cursor()
    phase_name, target_tps, workers, rpc = get_current_tps_phase(elapsed)
    for i, (_, query, headers) in enumerate(QUERIES):
        try:
            cur.execute(query)
            rows = cur.fetchall()
            for row in rows:
                processed = []; first = True
                for val in row:
                    if first:
                        processed.append(val.strftime("%Y-%m-%d %H:%M:%S") if isinstance(val, datetime) else str(val))
                        processed.append(phase_name); processed.append(target_tps)
                        processed.append(workers); processed.append(rpc)
                        first = False; continue
                    if val is None: processed.append("")
                    elif isinstance(val, bool): processed.append("Yes" if val else "No")
                    elif isinstance(val, datetime): processed.append(val.strftime("%Y-%m-%d %H:%M:%S"))
                    else: processed.append(val)
                sheet_data[i].append(processed)
            if not rows:
                r = [datetime.now().strftime("%Y-%m-%d %H:%M:%S"), phase_name, target_tps, workers, rpc, "NO DATA"]
                r.extend([""] * (len(headers) - 6)); sheet_data[i].append(r)
        except Exception as e:
            r = [datetime.now().strftime("%Y-%m-%d %H:%M:%S"), phase_name, target_tps, workers, rpc, f"ERROR: {e}"]
            r.extend([""] * (len(headers) - 6)); sheet_data[i].append(r); conn.rollback()
    conn.commit()
    return phase_name, target_tps

def save_workbook(wb, sheet_data, path):
    for i, (name, _, _) in enumerate(QUERIES):
        ws = wb[name]
        for row in sheet_data[i]: ws.append(row)
        auto_fit_columns(ws)
        if ws.max_row > 1: ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    wb.save(path)

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--duration", type=int, default=DURATION_SECONDS)
    parser.add_argument("--interval", type=int, default=SAMPLE_INTERVAL)
    parser.add_argument("--output-dir", type=str, default=OUTPUT_DIR)
    parser.add_argument("--host", type=str, default=DB_CONFIG["host"])
    parser.add_argument("--port", type=int, default=DB_CONFIG["port"])
    args = parser.parse_args()
    DB_CONFIG["host"] = args.host; DB_CONFIG["port"] = args.port

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    outfile = f"publisher_metrics_txnsize_{ts}.xlsx"
    outpath = os.path.join(args.output_dir, outfile)

    print("=" * 65)
    print("  Publisher Monitor — Transaction Size Experiment")
    print("=" * 65)
    print(f"  Host: {DB_CONFIG['host']}:{DB_CONFIG['port']} | DB: {DB_CONFIG['dbname']}")
    print(f"  Interval: {args.interval}s | Duration: {args.duration}s | Output: {outpath}")
    cumulative = 0
    for name, tps, w, rpc, dur in TPS_SCHEDULE:
        m1, s1 = divmod(cumulative, 60); m2, s2 = divmod(cumulative + dur, 60)
        print(f"    {m1}:{s1:02d}-{m2}:{s2:02d} -> {tps:>5,} TPS | {w:>2} workers | {rpc} rows/commit")
        cumulative += dur
    print("=" * 65)

    print("\n  Connecting..."); conn = psycopg2.connect(**DB_CONFIG); conn.autocommit = False; print("  Connected.")
    wb = create_workbook(); sheet_data = [[] for _ in QUERIES]
    print(f"  Monitoring started at {datetime.now().strftime('%H:%M:%S')}...\n")

    start = time.time(); end = start + args.duration; count = 0
    try:
        while time.time() < end:
            ss = time.time(); elapsed = time.time() - start
            pn, tt = collect_sample(conn, sheet_data, elapsed); count += 1
            try:
                cur = conn.cursor()
                cur.execute("SELECT pg_wal_lsn_diff(sent_lsn, replay_lsn) FROM pg_stat_replication LIMIT 1")
                r = cur.fetchone(); lag = f"{r[0]:,} bytes" if r else "No subs"; conn.commit()
            except Exception: lag = "N/A"; conn.rollback()
            print(f"  {datetime.now().strftime('%H:%M:%S')} | Sample {count:>4} | {int(elapsed):>4}s | {pn:45s} | Lag: {lag:>15s} | Left: {max(0, int(end - time.time()))}s")
            sl = max(0, args.interval - (time.time() - ss))
            if sl > 0: time.sleep(sl)
    except KeyboardInterrupt: print("\n  Stopping...")
    finally:
        print(f"\n  Saving {outfile}..."); save_workbook(wb, sheet_data, outpath); conn.close()
        tr = sum(len(s) for s in sheet_data)
        print(f"  Done: {count} samples, {tr} rows, {os.path.getsize(outpath):,} bytes")

if __name__ == "__main__":
    main()
